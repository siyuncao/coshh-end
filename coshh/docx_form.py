"""
Fill the Oxford *COSHH Form, Chemistry Teaching Laboratory* `.docx`.

This module is the **writing layer** and nothing else. It takes an assessment
that somebody else has already reasoned out (see ``coshh.rules``) and stamps it
onto a copy of the user's own template. It does no network I/O, it decides no
chemistry, and it never invents a hazard.

Why it is written against raw XML rather than python-docx's friendly API
-----------------------------------------------------------------------
The template's checkboxes are Word **content controls** (``w:sdt`` carrying a
``w14:checkbox``), not plain characters. The tick you see is a literal glyph
(``☒`` U+2612 / ``☐`` U+2610) inside ``w:sdtContent``, but its authoritative
state lives in ``w:sdtPr/w14:checkbox/w14:checked``. Writing only the glyph
leaves the two disagreeing and Word re-renders the box from the control, i.e.
the tick silently vanishes on the form somebody signs. :func:`_set_box` writes
both halves, and :meth:`CoshhForm.save` refuses to write a file where any box
disagrees with its own glyph.

Two further traps, both measured on the real template:

1. ``python-docx`` under-reports cells. Some cells are wrapped as
   ``w:sdt > w:sdtContent > w:tc`` (cell-level content controls), and
   ``row.cells`` only walks *direct* ``w:tc`` children of ``w:tr`` — the
   substance table's example row reports 2 cells for a 5-column row. Use
   :func:`_cells`.
2. Never call lxml's ``itertext()`` on python-docx elements: ``w:tc``, ``w:p``
   and ``w:r`` each expose a python-docx ``.text`` property that ``itertext()``
   picks up, so ``"Title:"`` reads back as ``"Title:Title:Title:"``. Use
   :func:`_text`.

The full anatomy, with the evidence behind every claim, is in
``docs/template-anatomy.md``.

Safety posture
--------------
The tool drafts; a competent person checks and signs.

* ``Approved By`` and its ``Date`` are **never** written, and :meth:`save`
  asserts they are still empty.
* A substance with no classification does not get a tidy empty cell. It gets
  :data:`UNASSESSED_TEXT` in the Hazards column, where it is impossible to miss.
* An unrecognised exposure route or control measure raises ``ValueError``
  rather than quietly ticking nothing.
"""

from __future__ import annotations

import copy
import random
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Sequence, Tuple

import docx
from docx.text.paragraph import Paragraph

# --------------------------------------------------------------------------
# Namespaces and template geography
# --------------------------------------------------------------------------

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
W14 = "{http://schemas.microsoft.com/office/word/2010/wordml}"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

CHECKED = "☒"      # ☒ BALLOT BOX WITH X
UNCHECKED = "☐"    # ☐ BALLOT BOX

TABLE_HEADER = 0
TABLE_SCHEME = 1
TABLE_SUBSTANCES = 2
TABLE_RISK = 3
TABLE_WASTE = 4
TABLE_APPROVAL = 5

SUBSTANCE_HEADER_ROW = 0
EXAMPLE_ROW = 1
TEMPLATE_BLANK_ROWS = 4     # the template ships four blank substance rows

COL_NAME = 0
COL_AMOUNT = 1
COL_HAZARDS = 2
COL_ROUTE = 3
COL_CONTROL = 4

EXPOSURE_ROUTES: Tuple[str, ...] = ("Eyes", "Skin", "Inhalation", "Ingestion")

#: The eleven Control Measures options, in the template's own order and wording.
CONTROL_MEASURES: Tuple[str, ...] = (
    "In case of spill, consult a demonstrator, technician or senior member of staff",
    "Safety spectacles",
    "Lab coat",
    "Gloves",
    "Fumehood",
    "Keep away from naked flames and sources of ignition",
    "Heat using temperature-controlled water bath",
    "Not to be used if pregnant",
    "Do not store or use near water (store in oil)",
    "Add dropwise to solution",
    "Do not expose to air",
)

#: Pre-ticked on every blank row in the template — the department's standing
#: rule. Cloning preserves it; nothing here unticks it on its own.
ALWAYS_ON: Tuple[str, ...] = CONTROL_MEASURES[:3]

#: Risk implication row -> row index in table 3. The label sits in a *different
#: cell* from its checkbox, so paragraph-local label lookup is useless here.
RISK_ROWS: Dict[str, int] = {
    "Fire or Explosion": 1,
    "Thermal Runaway": 2,
    "Gas Release": 3,
    "Malodorous Substances": 4,
    "Special measures": 5,
}
RISK_COL_BOX = 1
RISK_COL_MEASURES = 2

#: Waste stream -> (row, column of its checkbox) in table 4. The label lives in
#: the next column along, padded with trailing spaces, so address by position.
WASTE_CELLS: Dict[str, Tuple[int, int]] = {
    "Halogenated": (1, 0),
    "Aqueous": (1, 2),
    "Hydrocarbon": (2, 0),
    "Named Waste": (2, 2),
    "Contaminated solid waste": (3, 0),
    "Silica/TLC": (3, 2),
}
WASTE_NOTE_CELL = (0, 1)

#: What goes in the Hazards column when no classification could be found.
#: Loud on purpose: a blank cell reads as "assessed, nothing to declare".
UNASSESSED_TEXT = "NO CLASSIFICATION FOUND - check the supplier's safety data sheet"

#: The template the user supplies. Git-ignored; see README.
DEFAULT_TEMPLATE = Path(__file__).resolve().parent.parent / "templates" / "coshh-template.docx"


# --------------------------------------------------------------------------
# Low-level XML helpers
# --------------------------------------------------------------------------

def _cells(tr) -> List:
    """Logical cells of a ``w:tr``, in order.

    Unwraps cell-level content controls (``w:sdt > w:sdtContent > w:tc``).
    Use this instead of ``python-docx``'s ``row.cells``, which misses them.
    """
    out = []
    for child in tr:
        if child.tag == W + "tc":
            out.append(child)
        elif child.tag == W + "sdt":
            for tc in child.iter(W + "tc"):
                out.append(tc)
                break
    return out


def _text(el) -> str:
    """All text under ``el``. Never uses lxml ``itertext()`` — see module docstring."""
    return "".join(t.text or "" for t in el.iter(W + "t"))


def _checkboxes(el) -> List:
    """Every checkbox content control under ``el``, in document order."""
    return [s for s in el.iter(W + "sdt")
            if s.find(W + "sdtPr/" + W14 + "checkbox") is not None]


def _is_checked(sdt) -> bool:
    checked = sdt.find(W + "sdtPr/" + W14 + "checkbox/" + W14 + "checked")
    return checked is not None and checked.get(W14 + "val") in ("1", "true")


def _box_label(sdt) -> str:
    """The option text belonging to a checkbox.

    Exposure Route puts the label *before* the box; Control Measures puts it
    *after*. Joining the runs on both sides within the box's own paragraph
    handles both, survives runs split mid-label, and cannot merge two options
    because each option is its own paragraph with exactly one box.

    Returns ``""`` when the label lives in another cell (risk / waste tables).
    """
    p = sdt.getparent()
    if p is None or p.tag != W + "p":
        return ""
    before, after, seen = [], [], False
    for child in p:
        if child is sdt:
            seen = True
            continue
        if child.tag == W + "r":
            (after if seen else before).append(_text(child))
    return ("".join(before) + "".join(after)).strip()


def _set_box(sdt, on: bool) -> None:
    """Tick or untick a checkbox, writing **both** halves of its state.

    The codepoint and font are read from ``w14:checkedState`` /
    ``w14:uncheckedState`` on the element itself rather than hardcoded, and
    ``rFonts`` ascii/hAnsi/eastAsia are rewritten because Word rewrites the
    font when a human toggles a box (5 of the template's 86 had drifted to
    ``Segoe UI Symbol``).
    """
    checkbox = sdt.find(W + "sdtPr/" + W14 + "checkbox")
    if checkbox is None:
        raise ValueError("not a checkbox content control")
    state = checkbox.find(W14 + "checkedState" if on else W14 + "uncheckedState")
    val, font = state.get(W14 + "val"), state.get(W14 + "font")

    checkbox.find(W14 + "checked").set(W14 + "val", "1" if on else "0")

    run = sdt.find(W + "sdtContent/" + W + "r")
    run.find(W + "t").text = chr(int(val, 16))

    rPr = run.find(W + "rPr")
    if rPr is None:
        rPr = run.makeelement(W + "rPr", {})
        run.insert(0, rPr)
    rFonts = rPr.find(W + "rFonts")
    if rFonts is None:
        rFonts = rPr.makeelement(W + "rFonts", {})
        rPr.insert(0, rFonts)
    for attr in ("ascii", "hAnsi", "eastAsia"):
        rFonts.set(W + attr, font)


def _unwrap_cell(tc):
    """Replace a cell-level ``w:sdt`` wrapper with its bare ``w:tc``.

    Placeholder cells (``showingPlcHdr``) keep printing their grey example text
    unless the control is removed, so writing to one means unwrapping it first.
    """
    parent = tc.getparent()
    if parent is not None and parent.tag == W + "sdtContent":
        sdt = parent.getparent()
        sdt.addprevious(tc)
        sdt.getparent().remove(sdt)
    return tc


def _strip_placeholder_style(p) -> None:
    """Drop grey-italic placeholder formatting from a paragraph's ``pPr/rPr``."""
    pPr = p.find(W + "pPr")
    if pPr is None:
        return
    rPr = pPr.find(W + "rPr")
    if rPr is None:
        return
    for tag in ("rStyle", "color", "i", "iCs"):
        el = rPr.find(W + tag)
        if el is not None:
            rPr.remove(el)


def _clear_cell(tc):
    """Unwrap ``tc``, reduce it to its first paragraph, and empty that paragraph.

    Returns the surviving paragraph. The template's value cells hold exactly one
    ``w:p`` with a ``w:pPr`` and no runs, so keeping that paragraph (rather than
    adding one) inherits the cell's spacing and style.
    """
    tc = _unwrap_cell(tc)
    paragraphs = tc.findall(W + "p")
    if not paragraphs:
        p = tc.makeelement(W + "p", {})
        tc.append(p)
        return p
    for extra in paragraphs[1:]:
        tc.remove(extra)
    p = paragraphs[0]
    for r in p.findall(W + "r"):
        p.remove(r)
    _strip_placeholder_style(p)
    return p


def _append_run(p, text: str):
    r = p.makeelement(W + "r", {})
    t = r.makeelement(W + "t", {})
    t.text = text
    t.set(XML_SPACE, "preserve")
    r.append(t)
    p.append(r)
    return r


def _write_cell(tc, text: str) -> None:
    """Write plain text into a cell, one paragraph per newline."""
    p = _clear_cell(tc)
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        if i:
            nxt = copy.deepcopy(p)
            for r in nxt.findall(W + "r"):
                nxt.remove(r)
            p.addnext(nxt)
            p = nxt
        _append_run(p, line)


def _norm(s: str) -> str:
    """Loose key for matching a caller's label against the template's wording."""
    return " ".join(str(s).split()).rstrip(":").casefold()


# --------------------------------------------------------------------------
# The form
# --------------------------------------------------------------------------

@dataclass
class RenderReport:
    """What was written, and what a human still has to do."""

    out_path: Path
    substances: List[str] = field(default_factory=list)
    unassessed: List[str] = field(default_factory=list)
    needs_review: List[str] = field(default_factory=list)


class CoshhForm:
    """A COSHH form open for editing.

    Load the user's template, call the setters, then :meth:`save`. Nothing is
    written to disk until ``save``, and ``save`` validates before it writes.
    """

    def __init__(self, template_path=None):
        self.template_path = Path(template_path) if template_path else DEFAULT_TEMPLATE
        if not self.template_path.is_file():
            raise FileNotFoundError(
                "COSHH template not found at {}. Supply your own template "
                "with template_path=... (see README).".format(self.template_path)
            )
        self._doc = docx.Document(str(self.template_path))
        self._tables = self._doc.tables
        if len(self._tables) < 6:
            raise ValueError(
                "template has {} tables, expected 6 — is this the Oxford "
                "Chemistry Teaching Laboratory COSHH form?".format(len(self._tables))
            )
        # First data row of the substance table; moves if the example is dropped.
        self._data_row_start = EXAMPLE_ROW + 1
        self._example_dropped = False
        self._rows_prepared = False

    # -- generic access ---------------------------------------------------

    def _row(self, table_index: int, row_index: int):
        return self._tables[table_index]._tbl.findall(W + "tr")[row_index]

    def _cell(self, table_index: int, row_index: int, col_index: int):
        cells = _cells(self._row(table_index, row_index))
        try:
            return cells[col_index]
        except IndexError:
            raise IndexError(
                "table {} row {} has {} cells, no column {}".format(
                    table_index, row_index, len(cells), col_index)
            )

    def _used_sdt_ids(self) -> set:
        ids = set()
        for el in self._doc.element.body.iter(W + "id"):
            parent = el.getparent()
            if parent is not None and parent.tag == W + "sdtPr":
                ids.add(el.get(W + "val"))
        return ids

    # -- header ------------------------------------------------------------

    def set_header(self, title: str = "", name: str = "",
                   date: str = "", college: str = "") -> None:
        """Fill Title / Name / Date / College. Empty strings are skipped."""
        for value, (row, col) in (
            (title, (0, 1)),
            (name, (1, 1)),
            (date, (1, 3)),
            (college, (2, 1)),
        ):
            if value:
                _write_cell(self._cell(TABLE_HEADER, row, col), value)

    def set_year(self, year=None) -> None:
        """Underline 1, 2 or 3 in the ``1  /  2  /  3`` cell.

        The cell must still read ``1  /  2  /  3`` afterwards — a marked year is
        an underline, not a deletion. ``None`` leaves the cell untouched.
        """
        if year in (None, ""):
            return
        choice = str(year).strip()
        if choice not in ("1", "2", "3"):
            raise ValueError("year must be 1, 2 or 3 (got {!r})".format(year))

        tc = self._cell(TABLE_HEADER, 2, 3)
        p = next(iter(tc.iter(W + "p")))
        runs = p.findall(W + "r")
        if not runs:
            raise ValueError("Year cell has no runs; template differs from expectation")
        rPr = runs[0].find(W + "rPr")
        for r in runs:
            p.remove(r)
        for piece in ("1", "  /  ", "2", "  /  ", "3"):
            r = p.makeelement(W + "r", {})
            if rPr is not None:
                new_rPr = copy.deepcopy(rPr)
                r.append(new_rPr)
                if piece == choice:
                    u = new_rPr.makeelement(W + "u", {})
                    u.set(W + "val", "single")
                    new_rPr.append(u)
            t = r.makeelement(W + "t", {})
            t.text = piece
            t.set(XML_SPACE, "preserve")
            r.append(t)
            p.append(r)

    # -- reaction scheme ---------------------------------------------------

    def set_reaction_scheme(self, text: str = "", image_path=None) -> None:
        """Write the Reaction Scheme box.

        The cell is a placeholder content control; it is unwrapped so the grey
        example text does not survive into the printed form. ``image_path``
        drops in a picture (a ChemDraw export, say) after any text.
        """
        if not text and not image_path:
            return
        tc = self._cell(TABLE_SCHEME, 1, 0)
        p = _clear_cell(tc)
        if text:
            lines = str(text).split("\n")
            for i, line in enumerate(lines):
                if i:
                    nxt = copy.deepcopy(p)
                    for r in nxt.findall(W + "r"):
                        nxt.remove(r)
                    p.addnext(nxt)
                    p = nxt
                _append_run(p, line)
        if image_path:
            path = Path(image_path)
            if not path.is_file():
                raise FileNotFoundError("reaction scheme image not found: {}".format(path))
            if text:
                nxt = copy.deepcopy(p)
                for r in nxt.findall(W + "r"):
                    nxt.remove(r)
                p.addnext(nxt)
                p = nxt
            Paragraph(p, self._doc).add_run().add_picture(str(path))

    # -- substance table ---------------------------------------------------

    @property
    def data_row_count(self) -> int:
        """How many substance rows are available to fill."""
        rows = self._tables[TABLE_SUBSTANCES]._tbl.findall(W + "tr")
        return len(rows) - self._data_row_start

    def prepare_substance_rows(self, n: int, keep_example: bool = False) -> None:
        """Make room for ``n`` substances.

        The worked example (``e.g. 3 M HNO3``) is deleted unless ``keep_example``:
        its cells are placeholder content controls, and Word prints placeholder
        text, so a handed-in form would still carry the example. Extra rows are
        deep copies of the first genuinely blank row — never of the example row,
        which has demo ticks — with every ``w:sdtPr/w:id`` regenerated, because
        duplicate content-control ids are what makes Word complain on open.

        Never shrinks below the template's four blank rows.
        """
        if self._rows_prepared:
            raise RuntimeError("prepare_substance_rows() may only be called once")
        tbl = self._tables[TABLE_SUBSTANCES]._tbl

        if not keep_example:
            rows = tbl.findall(W + "tr")
            tbl.remove(rows[EXAMPLE_ROW])
            self._data_row_start = EXAMPLE_ROW
            self._example_dropped = True

        rows = tbl.findall(W + "tr")
        prototype = rows[self._data_row_start]     # a blank row, with the standing ticks
        last = rows[-1]
        have = len(rows) - self._data_row_start
        want = max(int(n), TEMPLATE_BLANK_ROWS)
        used = self._used_sdt_ids()

        for _ in range(want - have):
            clone = copy.deepcopy(prototype)
            for el in clone.iter(W + "id"):
                parent = el.getparent()
                if parent is not None and parent.tag == W + "sdtPr":
                    while True:
                        fresh = str(random.randint(1, 2 ** 31 - 1))
                        if fresh not in used:
                            break
                    used.add(fresh)
                    el.set(W + "val", fresh)
            last.addnext(clone)
            last = clone
        self._rows_prepared = True

    def _substance_row(self, i: int):
        rows = self._tables[TABLE_SUBSTANCES]._tbl.findall(W + "tr")
        index = self._data_row_start + i
        if not (self._data_row_start <= index < len(rows)):
            raise IndexError(
                "no substance row {} — {} data rows available; call "
                "prepare_substance_rows() first".format(i, self.data_row_count)
            )
        return rows[index]

    def _option_boxes(self, tc) -> "Dict[str, object]":
        """label -> checkbox, for a cell whose options carry their own labels."""
        out = {}
        for box in _checkboxes(tc):
            label = _box_label(box)
            if label:
                out[_norm(label)] = box
        return out

    def fill_substance(self, i: int, name: str, amount: str = "",
                       hazards: Sequence = (), routes: Sequence[str] = (),
                       controls: Sequence[str] = (),
                       controls_off: Sequence[str] = (),
                       unassessed: bool = False, note: str = "") -> None:
        """Fill one substance row. ``i`` is 0-based over the *data* rows.

        ``hazards`` may be strings or ``{"code": ..., "text": ...}`` mappings;
        mappings are rendered as ``"H225 - Highly flammable liquid and vapour"``,
        one per line. ``unassessed=True`` puts :data:`UNASSESSED_TEXT` at the top
        of the Hazards cell instead of leaving it looking clean.

        Exposure routes are set exhaustively (an unlisted route is explicitly
        unticked, because "no route ticked" is a claim the form makes). Control
        measures are **additive**: the template's standing spill / spectacles /
        lab coat ticks survive unless named in ``controls_off``.

        An unrecognised route or control raises ``ValueError`` — a silent no-op
        would mean a control measure the chemist believes is on the form.
        """
        tr = self._substance_row(i)
        cells = _cells(tr)

        _write_cell(cells[COL_NAME], name)
        if amount:
            _write_cell(cells[COL_AMOUNT], amount)

        lines: List[str] = []
        if unassessed:
            lines.append(UNASSESSED_TEXT)
        for h in hazards:
            if isinstance(h, dict):
                code = str(h.get("code", "")).strip()
                text = str(h.get("text", "")).strip()
                lines.append("{} - {}".format(code, text) if code and text else (code or text))
            else:
                lines.append(str(h))
        if note:
            lines.append(str(note))
        if lines:
            _write_cell(cells[COL_HAZARDS], "\n".join(x for x in lines if x))

        # Exposure routes: closed set, set every box.
        wanted_routes = {_norm(r) for r in routes}
        route_boxes = self._option_boxes(cells[COL_ROUTE])
        unknown = wanted_routes - set(route_boxes)
        if unknown:
            raise ValueError(
                "unknown exposure route(s) {} — the form offers {}".format(
                    sorted(unknown), list(EXPOSURE_ROUTES))
            )
        for key, box in route_boxes.items():
            _set_box(box, key in wanted_routes)

        # Control measures: additive.
        control_boxes = self._option_boxes(cells[COL_CONTROL])
        for wanted, state in [(controls, True), (controls_off, False)]:
            for label in wanted:
                box = control_boxes.get(_norm(label))
                if box is None:
                    raise ValueError(
                        "unknown control measure {!r} — the form offers {}".format(
                            label, list(CONTROL_MEASURES))
                    )
                _set_box(box, state)

    # -- risk implications and waste ---------------------------------------

    def set_risk(self, kind: str, yes: bool, measures: str = "") -> None:
        """Tick (or clear) one Specific Safety or Risk Implication row.

        There is one ``Y`` box per row, not a Y/N pair: "no" is an unticked box.
        Addressed by row index, because the label sits in a different cell from
        its checkbox.
        """
        key = _norm(kind)
        row = None
        for label, index in RISK_ROWS.items():
            if _norm(label) == key:
                row = index
                break
        if row is None:
            raise ValueError(
                "unknown risk implication {!r} — the form offers {}".format(
                    kind, list(RISK_ROWS))
            )
        boxes = _checkboxes(self._cell(TABLE_RISK, row, RISK_COL_BOX))
        if not boxes:
            raise ValueError("risk row {} has no checkbox".format(row))
        _set_box(boxes[0], bool(yes))
        if measures:
            _write_cell(self._cell(TABLE_RISK, row, RISK_COL_MEASURES), measures)

    def set_waste(self, kinds: Sequence[str] = (), note: str = "") -> None:
        """Tick waste streams, and optionally write the free-text waste note."""
        lookup = {_norm(k): v for k, v in WASTE_CELLS.items()}
        for kind in kinds:
            pos = lookup.get(_norm(kind))
            if pos is None:
                raise ValueError(
                    "unknown waste stream {!r} — the form offers {}".format(
                        kind, list(WASTE_CELLS))
                )
            row, col = pos
            boxes = _checkboxes(self._cell(TABLE_WASTE, row, col))
            if not boxes:
                raise ValueError("waste cell {} has no checkbox".format(pos))
            _set_box(boxes[0], True)
        if note:
            row, col = WASTE_NOTE_CELL
            _write_cell(self._cell(TABLE_WASTE, row, col), note)

    # -- validation and output ---------------------------------------------

    def validate(self) -> None:
        """Raise unless the document is structurally sound and unsigned.

        Checks the three things that go wrong silently: a checkbox whose glyph
        disagrees with its state (Word would re-render and drop the tick),
        duplicate content-control ids (Word complains on open), and a ``w:tc``
        orphaned outside ``w:tr``/``w:sdtContent`` (a botched unwrap). Then it
        checks the one thing that must never go right: ``Approved By`` filled in.
        """
        body = self._doc.element.body

        mismatched = []
        for sdt in _checkboxes(body):
            glyph = _text(sdt.find(W + "sdtContent"))
            if _is_checked(sdt) != (glyph == CHECKED):
                mismatched.append(_box_label(sdt) or "<unlabelled>")
        if mismatched:
            raise ValueError(
                "{} checkbox(es) disagree with their glyph: {}".format(
                    len(mismatched), mismatched[:5])
            )

        ids = []
        for el in body.iter(W + "id"):
            parent = el.getparent()
            if parent is not None and parent.tag == W + "sdtPr":
                ids.append(el.get(W + "val"))
        if len(ids) != len(set(ids)):
            dupes = sorted({i for i in ids if ids.count(i) > 1})
            raise ValueError("duplicate content-control ids: {}".format(dupes[:5]))

        stray = [tc for tc in body.iter(W + "tc")
                 if tc.getparent() is not None
                 and tc.getparent().tag not in (W + "tr", W + "sdtContent")]
        if stray:
            raise ValueError("{} table cell(s) outside w:tr/w:sdtContent".format(len(stray)))

        for col, what in ((1, "Approved By"), (3, "Approved Date")):
            if _text(self._cell(TABLE_APPROVAL, 0, col)).strip():
                raise ValueError(
                    "{} is filled in — only a competent person signs this form".format(what)
                )

    def save(self, path) -> Path:
        """Validate, then write the document. Returns the path written."""
        self.validate()
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self._doc.save(str(path))
        return path

    # -- read-back helpers (used by tests and by anyone auditing output) ----

    def header_values(self) -> Dict[str, str]:
        return {
            "title": _text(self._cell(TABLE_HEADER, 0, 1)),
            "name": _text(self._cell(TABLE_HEADER, 1, 1)),
            "date": _text(self._cell(TABLE_HEADER, 1, 3)),
            "college": _text(self._cell(TABLE_HEADER, 2, 1)),
            "year": _text(self._cell(TABLE_HEADER, 2, 3)),
        }

    def substance_row_values(self, i: int) -> Dict[str, object]:
        """Everything row ``i`` currently says — the shape tests assert on."""
        cells = _cells(self._substance_row(i))
        return {
            "name": _text(cells[COL_NAME]),
            "amount": _text(cells[COL_AMOUNT]),
            "hazards": _text(cells[COL_HAZARDS]),
            "routes": [_box_label(b) for b in _checkboxes(cells[COL_ROUTE])
                       if _is_checked(b)],
            "controls": [_box_label(b) for b in _checkboxes(cells[COL_CONTROL])
                         if _is_checked(b)],
        }

    def risk_values(self) -> Dict[str, Dict[str, object]]:
        return {
            label: {
                "yes": _is_checked(_checkboxes(self._cell(TABLE_RISK, row, RISK_COL_BOX))[0]),
                "measures": _text(self._cell(TABLE_RISK, row, RISK_COL_MEASURES)).strip(),
            }
            for label, row in RISK_ROWS.items()
        }

    def waste_values(self) -> List[str]:
        out = []
        for label, (row, col) in WASTE_CELLS.items():
            boxes = _checkboxes(self._cell(TABLE_WASTE, row, col))
            if boxes and _is_checked(boxes[0]):
                out.append(label)
        return out


def open_document(path, template_path=None) -> CoshhForm:
    """Re-open a generated file through :class:`CoshhForm`'s read-back helpers.

    The class is happy to load any file with the template's shape, so this is
    just ``CoshhForm(generated_file)`` with a name that says what it is for:
    auditing what actually landed in a document, including in tests.
    """
    form = CoshhForm(path)
    # A generated file has no worked-example row unless it was kept, and the
    # caller re-reading it wants row 0 to mean "first substance".
    rows = form._tables[TABLE_SUBSTANCES]._tbl.findall(W + "tr")
    if len(rows) > 1 and not _text(_cells(rows[1])[COL_NAME]).strip().lower().startswith("e.g."):
        form._data_row_start = EXAMPLE_ROW
    return form


# --------------------------------------------------------------------------
# Assessment dict -> document
# --------------------------------------------------------------------------

def render(assessment: Dict, out_path, template_path=None,
           keep_example: bool = False) -> RenderReport:
    """Turn an assessment dict into a filled COSHH form on disk.

    ::

        {
          "title": str, "name": str, "date": str, "college": str,
          "year": "1" | "2" | "3" | None,
          "scheme_note": str | None,
          "scheme_image": str | None,
          "substances": [
            {"name": str, "cas": str | None, "amount": str,
             "hazards": [{"code": "H225", "text": "..."}],
             "exposure": ["Eyes", "Skin"],
             "controls": ["Gloves", "Fumehood"],
             "unknown": bool, "note": str | None}
          ],
          "implications": {"Gas Release": {"yes": True, "prevention": "..."}},
          "waste": ["Aqueous"],
          "waste_note": str | None,
        }

    Nothing here decides chemistry: every tick comes from the dict. The returned
    :class:`RenderReport` lists what a human still has to settle — which always
    includes ``Approved By``, because this tool never signs.
    """
    substances = list(assessment.get("substances") or [])
    form = CoshhForm(template_path)
    form.prepare_substance_rows(len(substances), keep_example=keep_example)

    form.set_header(
        title=assessment.get("title") or "",
        name=assessment.get("name") or "",
        date=assessment.get("date") or "",
        college=assessment.get("college") or "",
    )
    form.set_year(assessment.get("year"))
    form.set_reaction_scheme(
        text=assessment.get("scheme_note") or "",
        image_path=assessment.get("scheme_image"),
    )

    report = RenderReport(out_path=Path(out_path))
    for i, s in enumerate(substances):
        name = s.get("name") or ""
        unknown = bool(s.get("unknown"))
        form.fill_substance(
            i,
            name=name,
            amount=s.get("amount") or "",
            hazards=s.get("hazards") or (),
            routes=s.get("exposure") or (),
            controls=s.get("controls") or (),
            controls_off=s.get("controls_off") or (),
            unassessed=unknown,
            note=s.get("note") or "",
        )
        report.substances.append(name)
        if unknown:
            report.unassessed.append(name)
            report.needs_review.append(
                "{}: no classification found — read the supplier's safety data "
                "sheet and complete the Hazards row by hand".format(name or "<unnamed>")
            )

    for kind, value in (assessment.get("implications") or {}).items():
        if isinstance(value, dict):
            form.set_risk(kind, bool(value.get("yes")), value.get("prevention") or "")
        else:
            form.set_risk(kind, bool(value))

    form.set_waste(assessment.get("waste") or (), assessment.get("waste_note") or "")

    if len(substances) < TEMPLATE_BLANK_ROWS:
        report.needs_review.append(
            "{} substance row(s) left blank — delete them or fill them in".format(
                TEMPLATE_BLANK_ROWS - len(substances))
        )
    report.needs_review.append(
        "Scale and containment: check the quantities against what you will "
        "actually use on the bench"
    )
    report.needs_review.append("Approved By / Date: left blank for a competent person to sign")

    form.save(out_path)
    return report


def document_parts(path) -> List[str]:
    """The OPC part names inside a .docx — a cheap structural smoke test."""
    with zipfile.ZipFile(str(path)) as zf:
        if zf.testzip() is not None:
            raise ValueError("corrupt zip: {}".format(path))
        return zf.namelist()
